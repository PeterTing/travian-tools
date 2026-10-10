#!/usr/bin/env python3
"""P0-10 開局攻略清單：把 Peter 的 Excel（opening.xlsx）轉成靜態資料檔.

用法（在 repo 根目錄）::

    python -m venv /tmp/opening-venv
    /tmp/opening-venv/bin/pip install -r scripts/requirements-opening.txt
    /tmp/opening-venv/bin/python scripts/convert_opening_checklist.py \
        /path/to/opening.xlsx

輸出 ``backend/data/static/opening_checklist.json``（後端驗證步驟、前端顯示都讀這份）。
Excel 本身不進 repo（公開 repo，內容是 Caím 的攻略）；輸出檔裡記了來源檔的 SHA-256。

規則：
- 結構和數字一律以英文工作表為準（``4P - Farm``、``3P - Sim``），數字不改。
  公式在這裡重算（英雄等級固定 0 ＝ Excel 預設；依英雄等級重算獎勵不在範圍內），
  並跟 Excel 存的計算結果比對，對不上就失敗。
- 中文：建築、等級、任務名稱用下面的對照表（台灣 Travian 用語，跟網站其他地方一致）；
  「為什麼」（General 欄）能對上翻譯工作表的就用翻譯工作表，經 OpenCC s2twp
  轉成繁體再套用語修正；翻譯工作表沒有、或內容跟英文不一致（數字不同、翻錯）的，
  用 ``MANUAL_ADVICE`` 的人工翻譯。
- 跟部族有關的欄位（開拓者花費、打野兵）放在每一步的 ``by_tribe``，步驟本身不變。
"""

from __future__ import annotations

import argparse
import ast
import hashlib
import json
import operator
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl
from opencc import OpenCC
from openpyxl.worksheet.formula import ArrayFormula

REPO = Path(__file__).resolve().parents[1]
OUTPUT = REPO / "backend" / "data" / "static" / "opening_checklist.json"

HERO_LEVEL = 0  # Excel 預設；依英雄等級重算獎勵不在 P0-10 範圍內

TRIBES = ["romans", "gauls", "teutons", "egyptians", "huns", "vikings", "spartans"]
# Excel「Other」表的部族名稱 → 網站的部族代碼（維京、斯巴達 Excel 沒有）
EXCEL_TRIBE = {
    "Roman": "romans",
    "Gaul": "gauls",
    "Teuton": "teutons",
    "Egyptian": "egyptians",
    "Hun": "huns",
}

# 打野兵：Excel 讓你自己選（Other!B39:B46）；網站依部族帶一個預設。
# 有騎兵可選的部族用騎兵（攻略建議 8 隻＝騎兵的建議數量 7～12），其餘用唯一選項。
FARM_UNIT_BY_TRIBE = {
    "romans": "Equites Imperatoris",
    "gauls": "Theutates Thunders",
    "teutons": "Clubswinger",
    "egyptians": "Slave Malitia",
    "huns": "Steppe Rider",
}
# Excel 原文拼錯的兵種英文名 → 官方拼法（backend/data/static/troops.json name_en）。
# Excel 裡的查表（unit_cost）還是用原文當 key，只有輸出給網站的英文改成官方拼法。
OFFICIAL_UNIT_EN = {"Slave Malitia": "Slave Militia"}


def unit_en(unit: str) -> str:
    return OFFICIAL_UNIT_EN.get(unit, unit)


UNIT_ZH = {
    "Equites Imperatoris": "帝國騎士",
    "Legionnaire": "古羅馬步兵",
    "Theutates Thunders": "雷法師",
    "Phalanx": "方陣兵",
    "Clubswinger": "棍棒兵",
    "Mercenary": "僱傭兵",
    "Steppe Rider": "草原騎士",
    "Slave Malitia": "奴隸民兵",
}

STRATEGIES = [
    {
        "id": "4p-farm",
        "name": "4P 農開",
        "sheet": "4P - Farm",
        "zh_sheet": "4趴农开-已翻译",
        "parties": 4,
    },
    {
        "id": "3p-sim",
        "name": "3P 兵開",
        "sheet": "3P - Sim",
        "zh_sheet": "3趴兵开-已翻译",
        "parties": 3,
    },
]

# ---------------------------------------------------------------- 用語對照

BUILDING_ZH = {
    "Main Building": "村莊大樓",
    "Woodcutters": "伐木場",
    "Clay pits": "泥坑",
    "Iron mines": "鐵礦場",
    "Croplands": "農場",
    "All woodcutters": "所有伐木場",
    "All clay pits": "所有泥坑",
    "All iron mines": "所有鐵礦場",
    "All croplands": "所有農場",
    "One cropland": "一塊農場",
    "Rally Point": "集結點",
    "Rallypoint": "集結點",
    "Granary": "穀倉",
    "Warehouse": "倉庫",
    "Marketplace": "市場",
    "Embassy": "大使館",
    "Cranny": "山洞",
    "Crannies": "山洞",
    "Cranny sims": "所有山洞",
    "City wall": "城牆",
    "Barracs": "兵營",
    "Academy": "研究院",
    "Smithy": "盔甲廠",
    "Stable": "馬廄",
    "Workshop": "工場",
    "Town hall": "城鎮廳",
    "Residence": "行宮",
    "Grain mill": "麵粉廠",
    # 任務（不是建築）
    "Even growth": "任務：均衡發展",
    "Complete economy": "任務：完整經濟",
    "Population": "任務：人口",
    "Culture point production": "任務：文明點產量",
}

_BUILDING_KEY = {k.lower(): k for k in BUILDING_ZH}


def canonical_building(name: str) -> str:
    """Excel 的大小寫、空白不一致（Main building／Main Building ）."""
    key = " ".join(name.split()).lower()
    if key not in _BUILDING_KEY:
        raise KeyError(name)
    return _BUILDING_KEY[key]


TASK_KIND = {
    "Even growth": "task",
    "Complete economy": "task",
    "Population": "task",
    "Culture point production": "task",
}


def target_zh(building_en: str, task: Any) -> str | None:
    """英文 task 欄 → 中文目標等級；看不懂就丟例外（不要默默放過）."""
    if task is None:
        return None
    if isinstance(task, int | float) and not isinstance(task, bool):
        n = int(task)
        if building_en == "Population":
            return f"人口達到 {n}"
        if building_en == "Culture point production":
            return f"每日文明點產量達到 {n}"
        raise ValueError(f"numeric task for {building_en!r}: {task!r}")
    text = str(task).strip()
    patterns: list[tuple[str, str]] = [
        (r"to (\d+)", "升到 {0}"),
        (r"1x to (\d+)", "一塊升到 {0}"),
        (r"all to (\d+)", "全部升到 {0}"),
        (r"All to (\d+)", "全部升到 {0}"),
        (r"1x of all to (\d+)", "每種資源田各一塊升到 {0}"),
        (r"all res to (\d+)", "所有資源田升到 {0}"),
        (r"(\d+)x to (\d+)", "{0} 個升到 {1}"),
    ]
    for pattern, template in patterns:
        m = re.fullmatch(pattern, text)
        if m:
            return template.format(*m.groups())
    raise ValueError(f"unknown task text {text!r} for {building_en!r}")


# ---------------------------------------------------------------- 為什麼

# 翻譯工作表的段落經 OpenCC 後再修的用語（台灣 Travian 說法、翻錯的詞）
TERM_FIXES = [
    ("救援物資生產", "資源產量"),
    ("回覆英雄的生命值", "回復英雄的生命值"),
    ("建議使用模擬檢視", "建議用模擬器看看"),
    ("最多 48小時開第一個Party", "最多 48 小時開第一場派對"),
    ("最多2小時", "最多 2 小時"),
    ("如果你打野打得太快，那麼你可能缺乏cp", "如果你推進得很快，可能會缺文明點"),
    ("可以較便宜獲得額外cp", "可以用比較少的資源換到文明點"),
    ("減少開趴時間", "縮短派對時間"),
    ("啟用25%資源加成", "啟用 25% 資源加成"),
    ("根據可以打的綠洲情況決定點多少級戰力", "依可以打的綠洲數量，決定戰鬥力要點幾點"),
    ("。 將剩餘", "。將剩餘"),
    ("不可能24小時", "不可能 24 小時"),
    ("之後考慮進行冒險 2。", "之後考慮去冒險 2（獎勵是資源）。"),
    ("0綠洲: ", "0 個綠洲："),
    ("1-5 綠洲: 8到12 點", "1～5 個綠洲：8 到 12 點"),
    ("6以上綠洲: 12點以上", "6 個以上綠洲：12 點以上"),
    ("最多2點", "最多 2 點"),
    ("24小時", "24 小時"),
    ("步行 1 小時", "走 1 小時"),
    ("路程需要2 小時", "路程要 2 小時"),
    ("清理掉那裡的動物", "清掉那裡的動物"),
]

# (策略, 英文工作表的列) → 翻譯工作表的儲存格；內容和英文一致才放這裡
ADVICE_FROM_ZH_SHEET = {
    ("4p-farm", 3): ("4趴农开-已翻译", "A3"),
    ("4p-farm", 6): ("4趴农开-已翻译", "A6"),
    ("4p-farm", 8): ("4趴农开-已翻译", "A8"),
    ("4p-farm", 22): ("4趴农开-已翻译", "A23"),
    ("4p-farm", 28): ("4趴农开-已翻译", "A29"),
    ("4p-farm", 71): ("4趴农开-已翻译", "A62"),
    ("4p-farm", 91): ("3趴兵开-已翻译", "A91"),
    ("3p-sim", 3): ("4趴农开-已翻译", "A6"),
    ("3p-sim", 7): ("4趴农开-已翻译", "A8"),
    ("3p-sim", 25): ("4趴农开-已翻译", "A29"),
    ("3p-sim", 45): ("4趴农开-已翻译", "A45"),
    ("3p-sim", 86): ("3趴兵开-已翻译", "A91"),
}

_PRO_TIP = (
    "進階：如果資源收入比需要的多，還有空間加快拓荒，可以考慮升級城鎮廳，"
    "縮短派對的時間，早點拿到派對的文明點。"
)
_TH_REMINDER = (
    "重要提醒：已經在排隊的派對，升級城鎮廳不會縮短它的時間。"
    "派對排隊中才提高文明點產量的話，增加的部分要等下一場派對開始才算。"
)

# 翻譯工作表沒有、或內容跟英文對不上的（數字不同、翻錯）：人工翻譯
MANUAL_ADVICE = {
    # 翻譯表的版本是「騎兵 10 到 20」，英文是 7 到 12，且多了一句帝國騎士
    ("4p-farm", 40): (
        "打野要做多少兵很難說。Excel 裡有一些簡單的試算，可以幫你判斷划不划算。\n"
        "建議數量：\n"
        "- 騎兵 7 到 12 隻\n"
        "- 步兵 30 到 60 隻\n"
        "做兵的時機也很重要：綠洲要等動物全部清掉後才會開始產出，"
        "所以不建議一開始就做打野兵。最有效率的數量看情況，依附近綠洲的實際手感決定。"
    ),
    ("4p-farm", 58): "重要提醒：雷法師、帝國騎士和草原騎士可以單獨派去沒有動物的綠洲！",
    ("4p-farm", 61): "目標：12／16 小時內全部升到 4",
    ("4p-farm", 63): (
        "人口和文明點的獎勵（Excel 裡標橘色的）要看馬廄升到 3 還是 5，"
        "請留意，不要太早把它算進來。"
    ),
    ("4p-farm", 72): (
        "這裡依 ROI 決定要升哪些資源田，或是多投資打野兵；別忘了你已經在打野，"
        "資源收入比較多。也許更該把重點放在文明點產量，因為要在 4 場派對內拓荒，"
        "文明點產量本來就得升。5 級資源田的 ROI：\n"
        "伐木場 67.4 小時\n"
        "泥坑 67.7 小時\n"
        "鐵礦場 82.5 小時\n"
        "農場 76.7 小時\n"
        "ROI 以英雄 6 級領到的獎勵計算。不建議升鐵礦場，埃及人尤其不要。"
    ),
    ("4p-farm", 94): _PRO_TIP,
    ("4p-farm", 103): _TH_REMINDER,
    ("3p-sim", 5): "英雄只去第一次冒險",
    # 翻譯表那句多了「確定能清掉動物」，是 4P 版本的說法
    ("3p-sim", 21): (
        "持續打綠洲，直到要走 1 小時以上；之後考慮去冒險 2（資源）。"
        "冒險大約能拿到 2300 資源。"
    ),
    # 翻譯表對應的那句是「最多 4 小時」，英文是 12 小時
    ("3p-sim", 50): "目標：12 小時內所有資源田升到 4。",
    ("3p-sim", 51): (
        "建議：農場升到 6 級，因為每塊會多 1 點文明點產量。之後文明點的花費會越來越高，"
        "可以把一塊農場看成 1250 資源的文明點投資，這樣農場本身只算 2000，"
        "ROI 大約 66 小時，加上文明點換算的成本低，很划算。"
    ),
    ("3p-sim", 59): (
        "選做：麵粉廠 1 級也是一個選項，搭配 6 級農場、算進任務獎勵的 ROI 是 63 小時。"
        "不過要拆掉一個山洞才有空地。"
    ),
    # 翻譯表對應的那段是「山洞升到 7、3 次活動約 60 小時」，跟英文不同
    ("3p-sim", 63): (
        "接下來是比較難的部分，請注意！\n"
        "如果你打野打得好，預計不到 48 小時就能開第一場派對，平常每小時產出的文明點會不夠，"
        "要靠提高文明點產量補回來。一定要算清楚你需要多少文明點、還差多少，"
        "再用下面的表補足。記得你有 3 場派對、至少 48 小時的產出可以補回少掉的文明點。"
    ),
    ("3p-sim", 78): "目標：72 小時內開第一場派對",
    ("3p-sim", 80): "開拓者的花費依部族不同，請在上方選部族。",
    ("3p-sim", 90): _PRO_TIP,
    ("3p-sim", 99): _TH_REMINDER,
}

# General 欄裡不是建議的格子（表頭）
ADVICE_SKIP = {("3p-sim", 89)}


# ---------------------------------------------------------------- 公式

_OPS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.USub: operator.neg,
    ast.UAdd: operator.pos,
}
_REF = re.compile(r"(?:(?P<sheet>[A-Za-z]+)!)?\$?(?P<col>[A-Z]{1,2})\$?(?P<row>\d+)")


def _eval_ast(node: ast.AST) -> float:
    if isinstance(node, ast.Expression):
        return _eval_ast(node.body)
    if isinstance(node, ast.Constant) and isinstance(node.value, int | float):
        return float(node.value)
    if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
        return _OPS[type(node.op)](_eval_ast(node.left), _eval_ast(node.right))
    if isinstance(node, ast.UnaryOp) and type(node.op) in _OPS:
        return _OPS[type(node.op)](_eval_ast(node.operand))
    raise ValueError(f"unsupported formula node {ast.dump(node)}")


class Evaluator:
    """只算四則運算和儲存格參照；有函式（IF、xlookup…）的格子由呼叫端處理."""

    def __init__(self, wb: openpyxl.Workbook, overrides: dict[tuple[str, str], Any]):
        self.wb = wb
        self.overrides = overrides

    def cell(self, sheet: str, ref: str) -> float | None:
        if (sheet, ref) in self.overrides:
            return self.overrides[(sheet, ref)]
        value = self.wb[sheet][ref].value
        return self.value(sheet, value)

    def value(self, sheet: str, value: Any) -> float | None:
        if value is None:
            return None
        if isinstance(value, bool):
            raise TypeError("boolean is not a number")
        if isinstance(value, int | float):
            return float(value)
        if isinstance(value, ArrayFormula):
            raise TypeError(f"array formula needs special handling: {value.text}")
        if isinstance(value, str) and value.startswith("="):
            expr = value[1:]
            if re.search(r"[A-Za-z]+\(", expr):
                raise ValueError(f"function in formula needs special handling: {value}")

            def repl(m: re.Match[str]) -> str:
                ref_sheet = m.group("sheet") or sheet
                v = self.cell(ref_sheet, f"{m.group('col')}{m.group('row')}")
                if v is None:
                    raise ValueError(f"empty cell referenced in {value}")
                return repr(v)

            return _eval_ast(ast.parse(_REF.sub(repl, expr), mode="eval"))
        raise ValueError(f"not a number: {value!r}")


def num(x: float | None) -> int | float | None:
    if x is None:
        return None
    r = round(x, 2)
    return int(r) if r == int(r) else r


# ---------------------------------------------------------------- Other 表的查表


def other_tables(wb: openpyxl.Workbook, ev: Evaluator) -> dict[str, Any]:
    ws = wb["Other"]

    def col_map(key_col: str, val_col: str, rows: range) -> dict[str, Any]:
        out = {}
        for r in rows:
            k = ws[f"{key_col}{r}"].value
            if k is None:
                continue
            raw = ws[f"{val_col}{r}"].value
            out[str(k).strip()] = raw
        return out

    unit_cost = {
        k: num(ev.value("Other", v))
        for k, v in col_map("B", "C", range(39, 47)).items()
    }
    research_cost = {
        k: num(ev.value("Other", v))
        for k, v in col_map("B", "C", range(49, 57)).items()
    }
    capacity = col_map("B", "C", range(59, 67))
    settlers: dict[str, int] = {}
    for r in range(69, 74):
        name = ws[f"B{r}"].value
        total = sum(int(ws.cell(r, c).value) for c in range(4, 8))  # D..G
        settlers[EXCEL_TRIBE[name]] = total
    stable_task = col_map("G", "H", range(39, 47))  # unit → "to 3"/"to 5"
    stable_cost = col_map("H", "I", range(39, 41))  # "to 5"/"to 3" → cost
    stable_cp = col_map("H", "I", range(49, 51))
    stable_pop = col_map("H", "I", range(51, 53))
    warehouse_unit = col_map("G", "H", range(54, 62))  # unit → Warehouse / N/A
    return {
        "unit_cost": unit_cost,
        "research_cost": research_cost,
        "capacity": capacity,
        "settlers": settlers,
        "stable_task": stable_task,
        "stable_cost": stable_cost,
        "stable_cp": stable_cp,
        "stable_pop": stable_pop,
        "warehouse_unit": warehouse_unit,
        "warehouse5": {  # Other!I54:L54（只有帝國騎士：倉庫先升到 5）
            "cost": ws["J54"].value,
            "cp": ws["K54"].value,
            "pop": ws["L54"].value,
        },
        "warehouse7_after5": {  # Other!I55:L55
            "cost": ws["J55"].value,
            "cp": ws["K55"].value,
            "pop": ws["L55"].value,
        },
    }


# ---------------------------------------------------------------- 主要轉換


def text_of(v: Any) -> str:
    if isinstance(v, ArrayFormula):
        return v.text
    return "" if v is None else str(v)


class Converter:
    def __init__(self, xlsx: Path):
        self.xlsx = xlsx
        self.wb = openpyxl.load_workbook(xlsx, data_only=False)
        self.cached = openpyxl.load_workbook(xlsx, data_only=True)
        self.cc = OpenCC("s2twp")
        self.problems: list[str] = []
        self.notes: list[str] = []
        overrides = {("Tasks", "I2"): float(HERO_LEVEL)}
        for s in STRATEGIES:
            overrides[(s["sheet"], "J1")] = float(HERO_LEVEL)
        self.ev = Evaluator(self.wb, overrides)
        self.other = other_tables(self.wb, self.ev)
        self.advice_used: set[tuple[str, int]] = set()

    # -- 為什麼
    def zh_advice(self, strategy_id: str, row: int, en: str) -> dict[str, str] | None:
        key = (strategy_id, row)
        if key in ADVICE_SKIP:
            return None
        self.advice_used.add(key)
        if key in ADVICE_FROM_ZH_SHEET:
            sheet, ref = ADVICE_FROM_ZH_SHEET[key]
            raw = str(self.wb[sheet][ref].value).strip()
            text = self.cc.convert(raw)
            for a, b in TERM_FIXES:
                text = text.replace(a, b)
            text = "\n".join(line.rstrip() for line in text.strip().splitlines())
            return {"zh": text, "en": en.strip(), "source": f"{sheet}!{ref}"}
        if key in MANUAL_ADVICE:
            return {"zh": MANUAL_ADVICE[key], "en": en.strip(), "source": "manual"}
        raise ValueError(f"no Chinese for advice {key}: {en[:60]!r}")

    def check_cached(self, sheet: str, ref: str, computed: Any) -> None:
        cached = self.cached[sheet][ref].value
        is_number = isinstance(cached, int | float) and not isinstance(cached, bool)
        if is_number and (
            computed is None or abs(float(cached) - float(computed)) > 0.01
        ):
            self.problems.append(
                f"{sheet}!{ref}: computed {computed!r} != Excel {cached!r}"
            )

    def numeric(self, sheet: str, ref: str) -> int | float | None:
        v = self.wb[sheet][ref].value
        if isinstance(v, str) and not v.startswith("="):
            return None  # 表頭文字
        result = num(self.ev.value(sheet, v))
        self.check_cached(sheet, ref, result)
        return result

    # -- 一份策略
    def strategy(self, spec: dict[str, Any]) -> dict[str, Any]:
        sid, sheet = spec["id"], spec["sheet"]
        ws = self.wb[sheet]
        steps: list[dict[str, Any]] = []
        pending_why: dict[str, str] | None = None
        extra_header_row: int | None = None
        extra_intro: dict[str, str] | None = None

        last_row = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 3).value)
        for r in range(3, last_row + 1):
            a = ws.cell(r, 1).value
            c = ws.cell(r, 3).value
            why = None
            if a is not None and str(a).strip():
                # 合併儲存格（A8:A21 之類）只在第一列有字：「為什麼」掛在那一步
                why = self.zh_advice(sid, r, str(a))
            # 「便宜文明點建築」表頭列：I='CP'、J='Pop'
            if ws.cell(r, 9).value == "CP" and ws.cell(r, 10).value == "Pop":
                extra_header_row = r
                if why:
                    pending_why = why
                continue
            if c is None:
                if why:
                    # 表格前的說明（4P A91／3P A86）：當成選做段落的說明
                    extra_intro = why
                continue
            if pending_why and not why:
                why, pending_why = pending_why, None
            step = self.step(sid, sheet, ws, r, why, extra=extra_header_row is not None)
            if step is not None:
                steps.append(step)

        sections = self.sections(sid, steps, extra_intro)
        return {
            "id": sid,
            "name": spec["name"],
            "sheet": sheet,
            "parties": spec["parties"],
            # 主進度只算必做的步驟；選做（便宜的文明點建築）另外算
            "required_steps": sum(
                len(sec["steps"]) for sec in sections if not sec["optional"]
            ),
            "optional_steps": sum(
                len(sec["steps"]) for sec in sections if sec["optional"]
            ),
            "sections": sections,
        }

    def step(
        self,
        sid: str,
        sheet: str,
        ws: Any,
        r: int,
        why: dict[str, str] | None,
        extra: bool,
    ) -> dict[str, Any] | None:
        c_raw = ws.cell(r, 3).value
        d = ws.cell(r, 4).value
        e = ws.cell(r, 5).value
        step: dict[str, Any] = {"id": f"r{r:03d}", "row": r}
        name_en = text_of(c_raw).strip()
        by_tribe: dict[str, dict[str, Any]] = {}

        if isinstance(c_raw, ArrayFormula) and "Other!G54" in c_raw.text:
            # 4P 58 列：帝國騎士才要先把倉庫升到 5，其他打野兵這列不用做
            w5 = self.other["warehouse5"]
            step.update(
                kind="farm_unit_only",
                building="倉庫",
                building_en="Warehouse",
                target="升到 5（只有帝國騎士需要）",
                target_en="to 5 (only applicable for EI)",
                tier=None,
                cost=None,
                cp=None,
                pop=None,
                skip=True,
            )
            for tribe in TRIBES:
                unit = FARM_UNIT_BY_TRIBE.get(tribe)
                if unit and self.other["warehouse_unit"].get(unit) == "Warehouse":
                    by_tribe[tribe] = {
                        "target": "升到 5",
                        "target_en": "to 5",
                        "cost": w5["cost"],
                        "cp": w5["cp"],
                        "pop": w5["pop"],
                        "skip": False,
                    }
            step["reward_res"] = step["reward_exp"] = None
        elif name_en == "Research":
            step.update(
                kind="research",
                building="研究打野兵",
                building_en="Research farm unit",
                target=None,
                target_en=None,
                tier=None,
                cost=None,
                cp=None,
                pop=None,
                reward_res=None,
                reward_exp=None,
            )
            for tribe in TRIBES:
                unit = FARM_UNIT_BY_TRIBE.get(tribe)
                if unit is None:
                    by_tribe[tribe] = {"missing": "farm_unit"}
                    continue
                cost = self.other["research_cost"].get(unit)
                by_tribe[tribe] = {
                    "building": f"研究{UNIT_ZH[unit]}",
                    "building_en": f"Research {unit_en(unit)}",
                    "cost": cost,
                    "target": None if cost else "基本兵種，不用研究",
                    "skip": not cost,
                }
        elif name_en == "Farm units":
            count = int(d)
            step.update(
                kind="farm_units",
                building="打野兵",
                building_en="Farm units",
                target=f"訓練 {count} 隻",
                target_en=f"{count} to make",
                tier=None,
                cost=None,
                cp=None,
                pop=None,
                reward_res=None,
                reward_exp=None,
                count=count,
            )
            for tribe in TRIBES:
                unit = FARM_UNIT_BY_TRIBE.get(tribe)
                if unit is None:
                    by_tribe[tribe] = {"missing": "farm_unit"}
                    continue
                by_tribe[tribe] = {
                    "building": f"打野兵（{UNIT_ZH[unit]}）",
                    "building_en": f"Farm units ({unit_en(unit)})",
                    "cost": num(self.other["unit_cost"][unit] * count),
                }
            if count == 0:
                step["target"] = "Excel 預設 0 隻，依 ROI 自己決定"
        elif name_en.startswith("Train settler"):
            count = int(e)
            step.update(
                kind="settlers",
                building="開拓者",
                building_en=name_en,
                target=f"訓練 {count} 個",
                target_en=str(count),
                tier=None,
                cost=None,
                cp=None,
                pop=None,
                reward_res=None,
                reward_exp=None,
                count=count,
            )
            for tribe in TRIBES:
                total = self.other["settlers"].get(tribe)
                by_tribe[tribe] = (
                    {"cost": total * count} if total else {"missing": "settler_cost"}
                )
            # Excel 存的是表頭 D1 選的那個部族的結果：拿來對我們的查表
            chosen = EXCEL_TRIBE.get(str(ws["D1"].value))
            if chosen and by_tribe[chosen].get("cost") is not None:
                self.check_cached(sheet, f"F{r}", by_tribe[chosen]["cost"])
        elif name_en.startswith("Party"):
            n = int(name_en.split()[1])
            step.update(
                kind="party",
                building=f"第 {n} 場派對",
                building_en=name_en,
                target=None,
                target_en=None,
                tier=None,
                cost=self.numeric(sheet, f"F{r}"),
                cp=None,
                pop=None,
                reward_res=self.numeric(sheet, f"G{r}"),
                reward_exp=self.numeric(sheet, f"H{r}"),
            )
        elif name_en.startswith("Don't forget"):
            step.update(
                kind="note",
                building="開新村前先留 3000 資源",
                building_en=name_en,
                target=None,
                target_en=None,
                tier=None,
                cost=self.numeric(sheet, f"F{r}"),
                cp=None,
                pop=None,
                reward_res=None,
                reward_exp=None,
            )
        else:
            try:
                name_en = canonical_building(name_en)
            except KeyError:
                raise ValueError(
                    f"{sheet}!C{r}: unknown building {name_en!r}"
                ) from None
            is_stable_variant = isinstance(e, ArrayFormula)
            is_wh_variant = (
                isinstance(ws.cell(r, 6).value, ArrayFormula) and not is_stable_variant
            )
            tier = d if isinstance(d, int) else None
            step.update(
                kind=TASK_KIND.get(name_en, "build"),
                building=BUILDING_ZH[name_en],
                building_en=name_en,
                tier=tier,
            )
            if is_stable_variant:
                # 4P 56 列：馬廄升到 3（一般）或 5（帝國騎士）
                o = self.other
                base_task = "to 3"
                step.update(
                    target=target_zh(name_en, base_task),
                    target_en=base_task,
                    cost=o["stable_cost"][base_task],
                    cp=o["stable_cp"][base_task],
                    pop=o["stable_pop"][base_task],
                )
                for tribe in TRIBES:
                    unit = FARM_UNIT_BY_TRIBE.get(tribe)
                    task = o["stable_task"].get(unit) if unit else None
                    if task and task != base_task:
                        by_tribe[tribe] = {
                            "target": target_zh(name_en, task),
                            "target_en": task,
                            "cost": o["stable_cost"][task],
                            "cp": o["stable_cp"][task],
                            "pop": o["stable_pop"][task],
                        }
            elif is_wh_variant:
                # 4P 69 列：倉庫升到 7；帝國騎士已經先升到 5，花費不同
                o = self.other
                f_formula = ws.cell(r, 6).value.text
                default_cost = int(re.search(r",(\d+)\)$", f_formula).group(1))
                i_formula = ws.cell(r, 9).value.text
                j_formula = ws.cell(r, 10).value.text
                step.update(
                    target=target_zh(name_en, e),
                    target_en=e,
                    cost=default_cost,
                    cp=int(re.search(r",(\d+)\)$", i_formula).group(1)),
                    pop=int(re.search(r",(\d+)\)$", j_formula).group(1)),
                )
                for tribe in TRIBES:
                    unit = FARM_UNIT_BY_TRIBE.get(tribe)
                    if unit and o["warehouse_unit"].get(unit) == "Warehouse":
                        w = o["warehouse7_after5"]
                        by_tribe[tribe] = {
                            "cost": w["cost"],
                            "cp": w["cp"],
                            "pop": w["pop"],
                        }
            else:
                step.update(
                    target=target_zh(name_en, e),
                    target_en=None if e is None else str(e).strip(),
                    cost=self.numeric(sheet, f"F{r}"),
                    cp=self.numeric(sheet, f"I{r}"),
                    pop=self.numeric(sheet, f"J{r}"),
                )
            if extra and ws.cell(r, 8).value == "Resc per 1 CP":
                step["reward_res"] = step["reward_exp"] = None
                step["res_per_cp"] = self.numeric(sheet, f"G{r}")
            else:
                step["reward_res"] = self.numeric(sheet, f"G{r}")
                step["reward_exp"] = self.numeric(sheet, f"H{r}")

        step["why"] = why
        if by_tribe:
            step["by_tribe"] = by_tribe
        step.setdefault("skip", False)
        if not step["skip"]:
            del step["skip"]
        step["_extra"] = extra
        return step

    def sections(
        self, sid: str, steps: list[dict[str, Any]], extra_intro: dict[str, str] | None
    ) -> list[dict[str, Any]]:
        """依任務等級分段：等級第一次出現 n 時開新的一段（保持 Excel 的順序）."""
        sections: list[dict[str, Any]] = []
        current_max = 0
        for step in steps:
            extra = step.pop("_extra")
            if extra:
                if not sections or sections[-1]["id"] != "extra-cp":
                    sections.append(
                        {
                            "id": "extra-cp",
                            "title": "選做：便宜的文明點建築",
                            "tier": None,
                            # 選做：不算進主進度（PM 決定），段落自己顯示完成數
                            "optional": True,
                            "intro": extra_intro,
                            "steps": [],
                        }
                    )
                sections[-1]["steps"].append(step)
                continue
            tier = step.get("tier")
            countable = step["kind"] in ("build", "task") and isinstance(tier, int)
            if not sections or (countable and tier > current_max):
                if countable:
                    current_max = tier
                sections.append(
                    {
                        "id": f"tier-{current_max}",
                        "title": f"任務等級 {current_max}",
                        "tier": current_max,
                        "optional": False,
                        "intro": None,
                        "steps": [],
                    }
                )
            sections[-1]["steps"].append(step)
        return sections

    # -- 參考：任務獎勵表
    def tasks_reference(self) -> list[dict[str, Any]]:
        ws = self.wb["Tasks"]
        group_zh = {
            "General": "一般",
            "Production ": "資源",
            "Production": "資源",
            "Infrastructuur": "建築",
        }
        task_zh = {
            "Culture point production": "文明點產量",
            "Population": "人口",
            "Crop": "農場",
            "Wood": "伐木場",
            "Clay": "泥坑",
            "Iron": "鐵礦場",
            "All crop": "所有農場",
            "All wood": "所有伐木場",
            "All clay": "所有泥坑",
            "All iron": "所有鐵礦場",
            "Even growth": "均衡發展",
            "All resources to lvl": "所有資源田",
            "Warehouse": "倉庫",
            "Granary": "穀倉",
            "Barracks": "兵營",
            "Stable": "馬廄",
            "Academy": "研究院",
            "Smithy": "盔甲廠",
            "Town Hall": "城鎮廳",
            "Workshop": "工場",
            "Main Building": "村莊大樓",
            "Main building": "村莊大樓",
            "Cranny": "山洞",
            "Market": "市場",
            "Embassy": "大使館",
            "Residence / Palace": "行宮／皇宮",
            "Wall": "城牆",
            "Rally point": "集結點",
            "Sawmill": "鋸木廠",
            "Brickyard": "磚廠",
            "Iron Foundry": "鋼鐵鑄造廠",
            "Grain mill": "麵粉廠",
            "Bakery": "麵包店",
            "Party": "派對",
        }

        def target(name: str, raw: Any) -> str:
            if isinstance(raw, str) and raw.startswith("="):
                raw = num(self.ev.value("Tasks", raw))
            if name in ("Culture point production",):
                return f"每日 {raw}"
            if name == "Population":
                return f"達到 {raw}"
            if name == "Party":
                return f"{raw} 場"
            s = str(raw)
            for pat, tpl in [
                (r"lvl (\d+)", "升到 {0}"),
                (r"1 van elk (\d+)", "每種各一塊升到 {0}"),
                (r"alle (\d+)", "全部升到 {0}"),
            ]:
                m = re.fullmatch(pat, s)
                if m:
                    return tpl.format(*m.groups())
            raise ValueError(f"Tasks: unknown target {raw!r} for {name!r}")

        def table(col0: int, first: int, last: int, title: str) -> dict[str, Any]:
            groups: list[dict[str, Any]] = []
            for r in range(first, last + 1):
                name = ws.cell(r, col0).value
                tier = ws.cell(r, col0 + 1).value
                if name is None:
                    continue
                if tier == "Tier":
                    groups.append({"title": group_zh[name], "rows": []})
                    continue
                if not groups:
                    groups.append({"title": "一般", "rows": []})
                name = str(name).strip() if str(name).strip() in task_zh else name
                key = str(name).strip()
                if key not in task_zh:
                    raise ValueError(f"Tasks: unknown task {name!r}")
                col = openpyxl.utils.get_column_letter
                groups[-1]["rows"].append(
                    {
                        "task": task_zh[key],
                        "tier": tier,
                        "target": target(key, ws.cell(r, col0 + 2).value),
                        "reward_res": self.numeric("Tasks", f"{col(col0 + 3)}{r}"),
                        "reward_exp": self.numeric("Tasks", f"{col(col0 + 4)}{r}"),
                    }
                )
            return {"title": title, "groups": groups}

        return [
            table(1, 4, ws.max_row, "首村的任務"),
            table(13, 4, ws.max_row, "新村的任務"),
            table(7, 9, ws.max_row, "帳號的任務"),
        ]

    # -- 參考：派對與文明點（Helpful stuff 的範例，只列數字不計算）
    def party_reference(self) -> dict[str, Any]:
        ws = self.wb["Helpful stuff"]
        label_res = {
            "Costs": "花費",
            "Task rewards": "任務獎勵",
            "Party 1": "第 1 場派對",
            "Party 2": "第 2 場派對",
            "Party 3": "第 3 場派對",
            "Party 4": "第 4 場派對",
        }
        label_cp = {
            "CP Prod": "文明點產量（每日）",
            "Town Hall lv": "城鎮廳等級",
            "Party 1 CP": "第 1 場派對的文明點",
            "Party 2 CP": "第 2 場派對的文明點",
            "Party 3 CP": "第 3 場派對的文明點",
            "Party 4 CP": "第 4 場派對的文明點",
        }
        phases = []
        for n, first in enumerate(range(7, 19, 3), start=1):
            rows = []
            for r in range(first, first + 3):
                rows.append(
                    {
                        "res_label": label_res[ws[f"A{r}"].value],
                        "resources": self.numeric("Helpful stuff", f"C{r}"),
                        "cp_label": label_cp[ws[f"E{r}"].value],
                        "cp_value": self.numeric("Helpful stuff", f"G{r}"),
                        "hours": self.numeric("Helpful stuff", f"H{r}"),
                        # Excel 用範例輸入算好的結果（照抄，不在網站上重算）
                        "cp_left": num(self.cached["Helpful stuff"][f"I{r}"].value),
                    }
                )
            phases.append({"title": f"到第 {n} 場派對", "rows": rows})
        return {
            "example_inputs": {
                "production_per_hour": ws["D3"].value,
                "farming_per_hour": ws["D4"].value,
                "cp_to_go": ws["I4"].value,
            },
            "phases": phases,
        }

    def run(self) -> dict[str, Any]:
        data = {
            "version": 1,
            "source": {
                "file": self.xlsx.name,
                "sha256": hashlib.sha256(self.xlsx.read_bytes()).hexdigest(),
                "sheets": {s["id"]: s["sheet"] for s in STRATEGIES},
                "hero_level": HERO_LEVEL,
                "generated_by": "scripts/convert_opening_checklist.py",
            },
            "tribe_data": {
                "settler_cost": {t: self.other["settlers"].get(t) for t in TRIBES},
                "farm_unit": {
                    t: (
                        {
                            "en": unit_en(FARM_UNIT_BY_TRIBE[t]),
                            "zh": UNIT_ZH[FARM_UNIT_BY_TRIBE[t]],
                            "cost": self.other["unit_cost"][FARM_UNIT_BY_TRIBE[t]],
                        }
                        if t in FARM_UNIT_BY_TRIBE
                        else None
                    )
                    for t in TRIBES
                },
            },
            "strategies": [self.strategy(s) for s in STRATEGIES],
            "reference": {
                "tasks": self.tasks_reference(),
                "party_cp": self.party_reference(),
            },
        }
        unused = (set(ADVICE_FROM_ZH_SHEET) | set(MANUAL_ADVICE)) - self.advice_used
        if unused:
            self.problems.append(f"advice mappings not used: {sorted(unused)}")
        return data


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("xlsx", type=Path, help="Peter 的 opening.xlsx")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args(argv)
    conv = Converter(args.xlsx)
    data = conv.run()
    if conv.problems:
        print("\n".join(conv.problems), file=sys.stderr)
        return 1
    text = json.dumps(data, ensure_ascii=False, indent=1) + "\n"
    if "種族" in text:
        raise SystemExit("用「部族」，不用「種族」")
    args.output.write_text(text, encoding="utf-8")
    for s in data["strategies"]:
        print(
            f"{s['id']}: {s['required_steps']} required + {s['optional_steps']} optional steps, "
            f"{len(s['sections'])} sections"
        )
    print(
        f"wrote {args.output.relative_to(REPO) if args.output.is_relative_to(REPO) else args.output}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
